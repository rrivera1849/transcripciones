"""Exporters: plain text, SRT subtitles, and Word (.docx)."""

from __future__ import annotations

from pathlib import Path

from . import Transcript
from .turns import Turn, fmt_ts, group_turns, speaker_names


def _label(turn: Turn, names: dict[str, str]) -> str:
    if turn.speaker is None:
        return ""
    return names.get(turn.speaker, turn.speaker)


def to_txt(transcript: Transcript, names: dict[str, str] | None = None, title: str | None = None) -> str:
    """`names` are overrides on top of the default "Hablante N" labels."""
    names = speaker_names(transcript, names)
    lines: list[str] = []
    if title:
        lines += [title, "=" * len(title), ""]
    for t in group_turns(transcript.segments):
        who = _label(t, names)
        head = f"[{fmt_ts(t.start)}] {who}:" if who else f"[{fmt_ts(t.start)}]"
        lines.append(f"{head} {t.text}")
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def to_srt(transcript: Transcript, names: dict[str, str] | None = None) -> str:
    """One cue per segment (not per turn) so subtitles stay short."""
    names = speaker_names(transcript, names)
    out: list[str] = []
    n = 0
    for seg in transcript.segments:
        text = " ".join(seg.text.split())
        if not text:
            continue
        n += 1
        who = names.get(seg.speaker, seg.speaker) if seg.speaker else None
        body = f"{who}: {text}" if who else text
        out += [str(n), f"{fmt_ts(seg.start, True)} --> {fmt_ts(seg.end, True)}", body, ""]
    return "\n".join(out)


def to_docx(
    transcript: Transcript,
    path: str | Path,
    names: dict[str, str] | None = None,
    title: str = "Transcripción",
    subtitle: str | None = None,
) -> Path:
    from docx import Document
    from docx.shared import Pt, RGBColor

    names = speaker_names(transcript, names)
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading(title, level=1)
    meta = []
    if subtitle:
        meta.append(subtitle)
    if transcript.duration:
        meta.append(f"Duración: {fmt_ts(transcript.duration)}")
    if meta:
        p = doc.add_paragraph(" · ".join(meta))
        p.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    for t in group_turns(transcript.segments):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        who = _label(t, names)
        if who:
            r = p.add_run(f"{who}  ")
            r.bold = True
        r = p.add_run(f"[{fmt_ts(t.start)}]  ")
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        r.font.size = Pt(9)
        p.add_run(t.text)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def write_all(transcript: Transcript, out_dir: str | Path, stem: str, names=None, title=None) -> dict[str, Path]:
    """Write .json, .txt, .srt and .docx next to each other; return the paths."""
    import json

    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = {
        "json": out_dir / f"{stem}.json",
        "txt": out_dir / f"{stem}.txt",
        "srt": out_dir / f"{stem}.srt",
        "docx": out_dir / f"{stem}.docx",
    }
    paths["json"].write_text(json.dumps(transcript.to_dict(), ensure_ascii=False, indent=1), encoding="utf-8")
    paths["txt"].write_text(to_txt(transcript, names, title), encoding="utf-8")
    paths["srt"].write_text(to_srt(transcript, names), encoding="utf-8")
    to_docx(transcript, paths["docx"], names, title=title or stem)
    return paths
